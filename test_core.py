import copy,io,json,sqlite3,tempfile,unittest
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import database as db
from metrics import calculate,aggregate,default_templates,validate_template
from imports import validate_rows,read_file,differences,csv_bytes
from reports import excel_report,pdf_report,analysis
from auth import hash_password,verify
E={'code':'tester-entry','role':'entry'};R={'code':'tester-review','role':'reviewer'};A={'code':'tester-admin','role':'admin'}

class PlatformTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.old=db.DB;db.DB=Path(self.tmp.name)/'test.db';db.init()
  self.template=next(t for t in db.templates() if t['id']=='shift-v1')
  self.rs=[r for r in db.results() if r['month']=='2026-09' and r['entity']=='SM-01']
 def tearDown(self):db.DB=self.old;self.tmp.cleanup()
 def test_formula(self):
  calc=calculate(self.rs,self.template)
  p=next(r for r in calc['rows'] if r['code']=='POST')
  self.assertAlmostEqual(p['score'],718/720*100);self.assertAlmostEqual(p['contribution'],718/720*15)
  self.assertTrue(calc['complete']);self.assertFalse(calc['final'])
 def test_weights_and_target(self):
  for t in default_templates():validate_template(t)
  t=copy.deepcopy(self.template);t['metrics'][0]['weight']=999
  with self.assertRaises(ValueError):validate_template(t)
  t=copy.deepcopy(self.template);t['metrics'][0]['target']=90
  self.assertEqual(calculate(self.rs,t)['rows'][0]['score'],100)
 def test_zero_missing_realzero(self):
  rs=copy.deepcopy(self.rs);rs[0]['actual']=0
  self.assertEqual(calculate(rs,self.template)['rows'][0]['score'],0)
  rs[0]['actual']=None
  self.assertIsNone(calculate(rs,self.template)['total'])
  rs[0]['actual']=0;rs[0]['plan']=0
  self.assertEqual(calculate(rs,self.template)['rows'][0]['state'],'zero_denominator')
 def test_na_rule(self):
  rs=copy.deepcopy(self.rs);rs[0].update(state='na',plan=0,actual=0)
  t=copy.deepcopy(self.template)
  self.assertIsNone(calculate(rs,t)['total'])
  t.update(approved=True,redistribute=True)
  calc=calculate(rs,t);self.assertTrue(calc['final']);self.assertAlmostEqual(sum(x['effective_weight'] for x in calc['rows']),100)
 def test_aggregate(self):
  t=copy.deepcopy(self.template);t['metrics']=[dict(t['metrics'][0],weight=100)]
  data=[dict(entity='A',metric='POST',state='ok',plan=10,actual=5),dict(entity='B',metric='POST',state='ok',plan=90,actual=90)]
  calc=calculate(aggregate(data,['A','B'],t),t)
  self.assertEqual(calc['rows'][0]['rate'],95)
  self.assertIsNone(calculate(aggregate(data,['A','B','C'],t),t)['total'])
 def test_import_idempotence_and_edit(self):
  start=len(db.results());n=db.import_rows(self.rs,E,db.digest(db.results()))
  self.assertEqual(len(db.results()),start)
  self.assertEqual(db.import_rows(self.rs,E,db.digest(db.results())),0)
  changed=copy.deepcopy(self.rs);changed[0]['actual']=700
  self.assertEqual(differences(changed,db.results())[0]['amal'],'Yangilash')
  db.import_rows(changed,E,db.digest(db.results()));self.assertEqual(len(db.results()),start)
  self.assertTrue(any(r['actual']==700 for r in db.results()))
 def test_concurrency_guard(self):
  with self.assertRaises(ValueError):db.import_rows(self.rs,E,'stale')
 def test_import_errors(self):
  row=copy.deepcopy(self.rs[0])
  for patch in [dict(plan=-1),dict(actual=99999),dict(plan=0,actual=0),dict(actual=None),dict(entity='unknown'),dict(month='2026-13'),dict(template='head-v1')]:
   self.assertTrue(validate_rows([{**row,**patch}])[1],patch)
  self.assertTrue(validate_rows([row,row])[1])
 def test_missing_and_na_valid(self):
  rs=[r for r in db.results() if r['month']=='2026-10']
  self.assertFalse(validate_rows(rs)[1])
 def test_roles(self):
  with self.assertRaises(PermissionError):db.import_rows(self.rs,R,db.digest(db.results()))
  with self.assertRaises(PermissionError):db.approve_template('shift-v1',E)
  with self.assertRaises(PermissionError):db.restore(json.loads(db.backup()),E)
 def test_snapshot_and_approval(self):
  db.approve_template('shift-v1',A);t=next(t for t in db.templates() if t['id']=='shift-v1');calc=calculate(self.rs,t)
  snap=dict(org='Test',month='2026-09',entity='SM-01',template=t,calculation=calc,analysis=analysis(calc),actions=[],inputs=self.rs)
  rid=db.create_report(snap,E);db.transition(rid,R);db.transition(rid,A)
  frozen=next(r for r in db.reports() if r['id']==rid)['payload']
  changed=copy.deepcopy(self.rs);changed[0]['actual']=600;db.import_rows(changed,E,db.digest(db.results()))
  self.assertEqual(next(r for r in db.reports() if r['id']==rid)['payload'],frozen)
  with self.assertRaises(ValueError):db.transition(rid,A)
 def test_draft_cannot_approve(self):
  c=calculate(self.rs,self.template);rid=db.create_report({'calculation':c},E);db.transition(rid,R)
  with self.assertRaises(ValueError):db.transition(rid,A)
 def test_restore_and_rollback(self):
  payload=json.loads(db.backup());count=len(db.results())
  db.restore(payload,A);self.assertEqual(len(db.results()),count)
  bad=copy.deepcopy(payload);r=json.loads(bad['tables']['results'][0]['payload']);r['entity']='NONEXISTENT';bad['tables']['results'][0]['payload']=json.dumps(r)
  with self.assertRaises(ValueError):db.restore(bad,A)
  self.assertEqual(len(db.results()),count)
 def test_password(self):
  h=hash_password('Test-secret-123');self.assertTrue(verify('Test-secret-123',h));self.assertFalse(verify('wrong',h))
 def test_exports(self):
  c=calculate(self.rs,self.template);s=dict(org='Namuna zavod',month='2026-09',entity='SM-01',template=self.template,calculation=c,analysis=analysis(c),actions=[],inputs=self.rs)
  rep=dict(id='TEST-001',payload=s,state='Qoralama',creator='tester',reviewer=None,approver=None)
  from openpyxl import load_workbook
  w=load_workbook(io.BytesIO(excel_report(rep)));self.assertIn('KPI_hisobi',w.sheetnames);self.assertEqual(w['KPI_hisobi'].max_row,11)
  pdf=pdf_report(rep);self.assertTrue(pdf.startswith(b'%PDF'));self.assertGreater(len(pdf),10000)
  root=Path(__file__).parents[1];qa=root.parent/'security_qa';qa.mkdir(exist_ok=True);(qa/'sample_report.pdf').write_bytes(pdf)
  content=csv_bytes([{'text':'=1+1'}]);self.assertIn(b"'=1+1",content)
 def test_sample_import(self):
  root=Path(__file__).parents[1]
  for name in ['Security_KPI_Namuna.csv','Security_KPI_Namuna.xlsx']:
   raw=read_file((root/'templates'/name).read_bytes(),name)
   clean,errors=validate_rows(raw.to_dict('records'));self.assertFalse(errors,errors);self.assertEqual(len(clean),60)

if __name__=='__main__':unittest.main(verbosity=2)
